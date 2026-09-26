# Copyright: (c) OpenSpug Organization. https://github.com/openspug/spug
# Copyright: (c) <spug.dev@gmail.com>
# Released under the AGPL-3.0 License.
# OpsForge：巡检产物——manifest JSON（供 diff）、HTML 报告、Excel 报告。
from django.conf import settings
from django.template import Context, Engine
from pathlib import Path
import json

from apps.ansible_ops.inspection import STATUS_CRIT, STATUS_WARN

# 与 executor.LOG_DIR 同款基路径约定（测试 settings 下自动落到工作区路径）
REPORTS_DIR = Path(settings.BASE_DIR) / 'storage' / 'reports'

# 运行时 v4.0.1 settings 未配置 TEMPLATES（上游后端零模板渲染），用独立 Engine
# 直接指到本 app 模板目录，不改 settings；autoescape 默认开启（主机名/说明均来自
# 采集数据，防 XSS），模板内不做任何 |safe
_ENGINE = Engine(dirs=[Path(__file__).resolve().parent.parent / 'templates'])

STATUS_LABEL = {'ok': '正常', 'warn': '警告', 'crit': '异常'}
CHECK_LABEL = {
    'cpu': 'CPU 压力', 'memory': '内存', 'disk': '磁盘',
    'process': '进程', 'port': '端口', 'cert': '证书到期',
}


def diff_results(prev_results, cur_results):
    """两轮结果 diff：逐 (主机, 检查) 比较 status 变化。

    prev/cur 结构同为 {host: {check: {'status': ...}}}；返回
    [{'host', 'check', 'from', 'to'}]，仅含状态发生变化或新增条目。
    """
    changes = []
    for host, checks in cur_results.items():
        prev_checks = (prev_results or {}).get(host, {})
        for check, info in checks.items():
            cur_status = info['status']
            prev_status = (prev_checks.get(check) or {}).get('status')
            if prev_status != cur_status:
                changes.append({
                    'host': host, 'check': check,
                    'from': prev_status or '-', 'to': cur_status,
                })
    return changes


def write_manifest(path, ts, results, changes):
    data = {'ts': ts, 'results': results, 'changes': changes}
    Path(path).write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding='utf-8')
    return data


def render_html(path, context):
    """渲染 HTML 报告（内联 CSS 单文件，离线可打开；异常项按 status 着色）。"""
    template = _ENGINE.get_template('inspection/report.html')
    Path(path).write_text(template.render(Context(context)), encoding='utf-8')


def _safe_cell(value):
    """防 Excel 公式注入：以 =+-@ 开头的字符串强制转文本（openpyxl 会把
    '=' 开头按公式写入，远端采集数据可构造 HYPERLINK 类恶意单元格）。"""
    if isinstance(value, str) and value[:1] in ('=', '+', '-', '@'):
        return "'" + value
    return value


def write_xlsx(path, ts, trigger_label, results, changes):
    """生成 Excel 汇总表；主表逐行着色（crit 红 / warn 橙），供人工留档翻阅。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    fills = {
        STATUS_CRIT: PatternFill('solid', fgColor='FFC7CE'),
        STATUS_WARN: PatternFill('solid', fgColor='FFEB9C'),
    }
    head_fill = PatternFill('solid', fgColor='4472C4')
    head_font = Font(color='FFFFFF', bold=True)

    wb = Workbook()
    ws = wb.active
    ws.title = '巡检报表'
    ws.append([_safe_cell(f'OpsForge 批量巡检报告（{ts}，{trigger_label}）')])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)
    ws.append([])
    ws.append(['主机', '检查项', '状态', '数值', '说明', '与上轮变化'])
    for cell in ws[3]:
        cell.fill = head_fill
        cell.font = head_font
    change_map = {(c['host'], c['check']): c for c in changes}
    for host, checks in results.items():
        for check, info in checks.items():
            change = change_map.get((host, check))
            change_text = '-' if not change else f"{change['from']} → {change['to']}"
            ws.append([
                _safe_cell(host), _safe_cell(CHECK_LABEL.get(check, check)),
                STATUS_LABEL.get(info['status'], info['status']),
                _safe_cell(info.get('value') or '-'),
                _safe_cell(info.get('message') or '-'),
                _safe_cell(change_text),
            ])
            row = ws.max_row
            fill = fills.get(info['status'])
            if fill:
                for cell in ws[row]:
                    cell.fill = fill
            ws.cell(row=row, column=3).alignment = Alignment(horizontal='center')
    for column, width in zip('ABCDEF', (12, 12, 8, 14, 48, 14)):
        ws.column_dimensions[column].width = width
    wb.save(path)
